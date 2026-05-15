"""Export pivot data with embedded charts to Excel (.xlsx)."""
from __future__ import annotations

from typing import List

try:
    import openpyxl
    from openpyxl.chart import BarChart, Reference
except ImportError as exc:  # pragma: no cover
    raise ImportError("openpyxl is required: pip install openpyxl") from exc

from macro_sync.schema import NutritionEntry
from macro_sync.exporters.pivot_summary import build_pivot, source_totals


_MACRO_COLS = ["calories", "protein_g", "carbs_g", "fat_g"]


def _write_pivot_sheet(ws, pivots) -> None:
    """Write daily pivot rows and embed a bar chart."""
    headers = ["date", "entries"] + _MACRO_COLS
    ws.append(headers)

    for day in pivots:
        ws.append([
            day.date.isoformat(),
            day.entry_count,
            day.total_calories,
            day.total_protein_g,
            day.total_carbs_g,
            day.total_fat_g,
        ])

    if len(pivots) == 0:
        return

    chart = BarChart()
    chart.type = "col"
    chart.title = "Daily Macros"
    chart.y_axis.title = "grams / kcal"
    chart.x_axis.title = "Date"
    chart.width = 20
    chart.height = 12

    n_rows = len(pivots) + 1  # +1 for header
    data_ref = Reference(ws, min_col=3, max_col=6, min_row=1, max_row=n_rows)
    cats_ref = Reference(ws, min_col=1, min_row=2, max_row=n_rows)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    ws.add_chart(chart, "H2")


def _write_source_sheet(ws, pivots) -> None:
    """Write per-source totals."""
    totals = source_totals(pivots)
    ws.append(["source", "calories", "protein_g", "carbs_g", "fat_g"])
    for src, bd in sorted(totals.items()):
        ws.append([src, bd.calories, bd.protein_g, bd.carbs_g, bd.fat_g])


def entries_to_pivot_chart_workbook(entries: List[NutritionEntry]) -> "openpyxl.Workbook":
    pivots = build_pivot(entries)
    wb = openpyxl.Workbook()
    ws_pivot = wb.active
    ws_pivot.title = "Daily Pivot"
    _write_pivot_sheet(ws_pivot, pivots)

    ws_src = wb.create_sheet("Source Totals")
    _write_source_sheet(ws_src, pivots)
    return wb


def entries_to_pivot_chart_bytes(entries: List[NutritionEntry]) -> bytes:
    import io
    wb = entries_to_pivot_chart_workbook(entries)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
